#!/usr/bin/env python3
"""Generate original ODF fixtures and convert with an isolated LibreOffice.

Requires python-pptx for the presentation source; ODT uses the standard library. Expected JSON is deliberately NOT generated here.
"""
from pathlib import Path
import argparse
import os
import signal
import subprocess
import tempfile
from xml.sax.saxutils import escape
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED, ZIP_STORED

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / 'tests' / 'fixtures'
NS = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
      'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
      'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
      'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
      'xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" '
      'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
      'xmlns:presentation="urn:oasis:names:tc:opendocument:xmlns:presentation:1.0" '
      'xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0"')


def p(text):
    return '<text:p>' + escape(text) + '</text:p>'


def cell(text, attrs=''):
    return f'<table:table-cell office:value-type="string" {attrs}>{p(text)}</table:table-cell>'


def row(cells):
    return '<table:table-row>' + ''.join(cells) + '</table:table-row>'


def table(rows):
    return ('<table:table table:name="SyntheticTable">'
            '<table:table-column table:number-columns-repeated="3"/>'
            + ''.join(rows) + '</table:table>')


def sources():
    ordinary = p('Synthetic document / Синтетический документ') + table([
        row([cell('Key'), cell('Value'), cell('Примечание')]),
        row([cell('alpha beta'), cell(''), cell('два слова')]),
        row([cell(''), cell('42'), cell('English and русский')]),
    ]) + p('After table / После таблицы')
    merged = p('Merged cells / Объединённые ячейки') + table([
        row([cell('Wide heading / Общий заголовок', 'table:number-columns-spanned="2"'),
             '<table:covered-table-cell/>', cell('Right')]),
        row([cell('Tall / Высокая', 'table:number-rows-spanned="2"'), cell('middle one'), cell('')]),
        row(['<table:covered-table-cell/>', cell('middle two'), cell('Конец')]),
    ]) + p('After merged table / После объединения')
    return [('ordinary_table.odt', 'text', ordinary), ('merged_table.odt', 'text', merged)]


def write_pptx(path):
    from pptx import Presentation
    from datetime import datetime
    deck = Presentation()
    deck.core_properties.author = 'legacy-office-parser synthetic fixtures'
    deck.core_properties.last_modified_by = 'legacy-office-parser synthetic fixtures'
    deck.core_properties.title = 'Original synthetic slide fixture'
    deck.core_properties.created = datetime(2026, 1, 1)
    deck.core_properties.modified = datetime(2026, 1, 1)
    first = deck.slides.add_slide(deck.slide_layouts[1])
    first.shapes.title.text = 'First title / Первый заголовок'
    first.placeholders[1].text = 'English body with two words\nРусский текст с пробелами'
    deck.slides.add_slide(deck.slide_layouts[6])
    third = deck.slides.add_slide(deck.slide_layouts[1])
    third.shapes.title.text = 'Third title / Третий заголовок'
    third.placeholders[1].text = 'Last body / Последний текст'
    third.notes_slide.notes_text_frame.text = 'Third slide notes only / Заметки только третьего слайда'
    # Exclude the template's slide-image/slide-number placeholders from notes.
    # LibreOffice exports the latter as a literal '*' in binary PPT.
    notes = third.notes_slide
    for shape in list(notes.shapes):
        if not shape.has_text_frame or shape.text_frame is None or shape.text != 'Third slide notes only / Заметки только третьего слайда':
            shape._element.getparent().remove(shape._element)
    deck.save(path)


def write_odf(path, kind, body):
    mime = f'application/vnd.oasis.opendocument.{kind}'
    content = (f'<?xml version="1.0" encoding="UTF-8"?><office:document-content {NS} office:version="1.2">'
               f'<office:body><office:{kind}>{body}</office:{kind}></office:body></office:document-content>')
    styles = (f'<?xml version="1.0" encoding="UTF-8"?><office:document-styles {NS} office:version="1.2">'
              '<office:styles><style:presentation-page-layout style:name="TitleBody">'
              '<presentation:placeholder presentation:object="title" svg:x="2cm" svg:y="1cm" svg:width="24cm" svg:height="4cm"/>'
              '<presentation:placeholder presentation:object="outline" svg:x="2cm" svg:y="6cm" svg:width="24cm" svg:height="4cm"/>'
              '</style:presentation-page-layout></office:styles><office:automatic-styles><style:page-layout style:name="Page">'
              '<style:page-layout-properties fo:page-width="28cm" fo:page-height="21cm" '
              'style:print-orientation="landscape"/></style:page-layout></office:automatic-styles>'
              '<office:master-styles><style:master-page style:name="Default" style:page-layout-name="Page"/>'
              '</office:master-styles></office:document-styles>')
    manifest = ('<?xml version="1.0" encoding="UTF-8"?><manifest:manifest '
                'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2">'
                f'<manifest:file-entry manifest:full-path="/" manifest:media-type="{mime}"/>'
                '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
                '<manifest:file-entry manifest:full-path="styles.xml" manifest:media-type="text/xml"/>'
                '</manifest:manifest>')
    with ZipFile(path, 'w') as archive:
        for name, data in [('mimetype', mime), ('content.xml', content), ('styles.xml', styles),
                           ('META-INF/manifest.xml', manifest)]:
            entry = ZipInfo(name, date_time=(2026, 1, 1, 0, 0, 0))
            entry.compress_type = ZIP_STORED if name == 'mimetype' else ZIP_DEFLATED
            archive.writestr(entry, data.encode('utf-8'))


def convert(lo, profile, source, destination, fmt, timeout):
    command = [lo, '-env:UserInstallation=' + profile.as_uri(), '--headless', '--nologo',
               '--nodefault', '--nofirststartwizard', '--convert-to', fmt,
               '--outdir', str(destination), str(source)]
    proc = subprocess.Popen(command, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                            text=True, start_new_session=True,
                            env={**os.environ, 'SAL_USE_VCLPLUGIN': 'svp',
                                 'DCONF_PROFILE': '/dev/null', 'GSETTINGS_BACKEND': 'memory'})
    try:
        output, _ = proc.communicate(timeout=timeout)
    except subprocess.TimeoutExpired:
        # Terminate only this invocation's process group, never other LO sessions.
        os.killpg(proc.pid, signal.SIGKILL)
        proc.communicate()
        raise RuntimeError(f'LibreOffice timed out after {timeout}s: {source.name}') from None
    if proc.returncode:
        raise RuntimeError(f'LibreOffice failed ({proc.returncode}): {output}')
    print(output.strip())
    target = destination / (source.stem + '.' + fmt.split(':')[0])
    if not target.is_file() or target.read_bytes()[:8] != bytes.fromhex('d0cf11e0a1b11ae1'):
        raise RuntimeError(f'LibreOffice did not produce a binary OLE file: {target}')
    return target


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--libreoffice', default='/usr/bin/libreoffice')
    parser.add_argument('--timeout', type=float, default=45)
    args = parser.parse_args()
    source_dir = FIXTURES / 'sources'
    source_dir.mkdir(parents=True, exist_ok=True)
    # Stage conversions to prevent stale binaries from satisfying validation.
    with tempfile.TemporaryDirectory(prefix='legacy-office-fixtures-', dir='/tmp') as temp:
        work = Path(temp)
        for name, kind, body in sources():
            source = source_dir / name
            write_odf(source, kind, body)
            fmt = 'doc:MS Word 97' if kind == 'text' else 'ppt:MS PowerPoint 97'
            target = convert(args.libreoffice, work / 'profile', source, work, fmt, args.timeout)
            (FIXTURES / target.name).write_bytes(target.read_bytes())
        source = source_dir / 'slides_notes.pptx'
        write_pptx(source)
        target = convert(args.libreoffice, work / 'profile', source, work,
                         'ppt:MS PowerPoint 97', args.timeout)
        (FIXTURES / target.name).write_bytes(target.read_bytes())
    print('Expected JSON was not modified.')


if __name__ == '__main__':
    main()
